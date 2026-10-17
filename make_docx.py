from docx import Document
d = Document()
d.add_heading("Foundations of Generative AI (ICT4403) - Assignment 2", 1)
rows = [("Student name", "Aditya Nigam"), ("Registration number", "230911366"),
        ("Assignment title", "Assignment 2 - Autoregressive Text Generation using RNN, BiRNN, LSTM and GRU"),
        ("System title", "Word-Level Autoregressive Text Generator for Alice's Adventures in Wonderland: SimpleRNN vs BiRNN vs LSTM vs GRU"),
        ("Dataset", "Alice's Adventures in Wonderland (Project Gutenberg #11)"),
        ("Dataset link", "https://www.gutenberg.org/cache/epub/11/pg11.txt"),
        ("GitHub repository", "https://github.com/aditya1nigam8/FGAI-Assignment2-Autoregressive-TextGen")]
t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
for a, b in rows:
    c = t.add_row().cells; c[0].text, c[1].text = a, b
d.save("AdityaNigam_230911366_FGAI_Assignment2.docx")
